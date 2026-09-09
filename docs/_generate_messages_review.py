"""One-off generator for User-Facing-Messages-Wording-Review.docx. Not part of the app."""
from __future__ import annotations

import os
from datetime import date

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT_PATH = os.path.join(
    os.path.dirname(__file__),
    "User-Facing-Messages-Wording-Review.docx",
)

HEADERS = [
    "Type",
    "Page / Feature",
    "Trigger / Condition",
    "Current Message",
    "Suggested Revision",
]


def set_run_font(run, size=10, bold=False, color=None):
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.bold = bold
    if color:
        run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, size=8, white=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    set_run_font(
        run,
        size=size,
        bold=bold,
        color=RGBColor(255, 255, 255) if white else None,
    )
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=8, white=True)
        shade_cell(table.rows[0].cells[i], "1F4E79")
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            text = val or ""
            set_cell_text(table.rows[r_idx + 1].cells[c_idx], text, size=8)
            if r_idx % 2 == 1:
                shade_cell(table.rows[r_idx + 1].cells[c_idx], "F2F2F2")
            if c_idx == 4 and any(
                flag in text
                for flag in ("FLAG", "DUPLICATE", "INCONSISTENT", "⚠")
            ):
                shade_cell(table.rows[r_idx + 1].cells[c_idx], "FFF2CC")
    doc.add_paragraph()
    return table


def R(type_, page, trigger, current, suggested):
    return [type_, page, trigger, current, suggested]


def bullet(doc, text, size=10):
    bp = doc.add_paragraph(style="List Bullet")
    run = bp.add_run(text)
    set_run_font(run, size=size)


def main():
    doc = Document()
    for section in doc.sections:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Cm(29.7)
        section.page_height = Cm(21.0)
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run(
        "Leave Forms — User-Facing Messages Inventory & Wording Review"
    )
    set_run_font(r, size=18, bold=True, color=RGBColor(31, 78, 121))

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        f"Generated from frontend source review · {date.today().isoformat()} · "
        "No code changes made"
    )
    set_run_font(r, size=10, color=RGBColor(89, 89, 89))

    doc.add_heading("1. Purpose", level=1)
    p = doc.add_paragraph()
    r = p.add_run(
        "This document inventories user-facing text currently hardcoded in the "
        "leaveforms-frontend app (toasts, banners, validations, confirmations, "
        "empty states, tooltips, error pages, helpers, and status labels). "
        "Suggested revisions are proposals only — for review and standardization "
        "before any code changes."
    )
    set_run_font(r, size=10)

    doc.add_heading("2. Scope notes", level=1)
    for n in [
        "No i18n/locale layer — copy is hardcoded in TS/TSX (centralized toast strings in src/constants/toastMessages.ts).",
        "Many API/Laravel validation errors are shown verbatim from the backend (not listed unless the frontend has a fallback string).",
        "Toast errors are not used — failures generally show as inline banners/alerts.",
        "Dynamic values ({name}, {type}, counts, dates) are shown with placeholders in braces.",
        "UI chrome labels (Submit, Cancel, table headers, nav items) are excluded unless instructional or status messaging.",
    ]:
        bullet(doc, n)

    doc.add_heading("3. Cross-cutting findings (review first)", level=1)
    findings = [
        (
            "Duplicate / near-duplicate meaning",
            "Same action uses different wording across pages, or identical wording for different contexts.",
        ),
        (
            "Inconsistent error tone",
            'Mix of “Unable to …” vs “Could not … Try again.” vs “… Please try again.”',
        ),
        (
            "Title casing inconsistency",
            'e.g. “Approve Leave Request” vs “Approve leave request”; “Confirm Distribution” vs “Confirm distribution”; filter “Pending conversion” vs badge “PENDING CONVERSION”.',
        ),
        (
            "Grammar issue",
            'Reminder toast: “pending leave request” (singular) when referring to multiple approvers/requests.',
        ),
        (
            "Technical / vague copy",
            'Login network error mentions “API … port 8000” (too technical). Empty states “No data to show” are generic.',
        ),
        (
            "Same title, different meaning",
            '“No remaining {type} credits” is reused for both a blocking banner and an LWOP informational banner.',
        ),
        (
            "Success toast pattern",
            'Most success toasts use “… successfully.” — keep this pattern; trim redundancy where the toast UI already implies success.',
        ),
    ]
    add_table(doc, ["Finding", "Detail"], findings)

    doc.add_heading("4. Recommended wording standards (proposed)", level=1)
    for s in [
        "Tone: clear, calm, direct; sentence case for titles except proper nouns / leave-type acronyms (VL, SL, EL, LWOP, SPL).",
        'Success toasts: “[Thing] [action] successfully.” Keep short.',
        'Error fallbacks: standardize on “Unable to [action]. Please try again.”',
        'Validation: “Please [action].” for missing required fields; “[Field] is required.” for free-text fields.',
        'Confirmations: “Are you sure you want to [action]?” for yes/no; reason modals: “Please provide a reason for [action].”',
        'Empty states: prefer specific (“No leave requests yet”) over generic (“No data to show”) when context is known.',
        "Status filters: Title Case consistently (Pending Approval, Pending Conversion). Badges may stay ALL CAPS as a visual system.",
    ]:
        bullet(doc, s)

    doc.add_heading("5. Inventory by type", level=1)

    # --- 5.1 Toasts ---
    doc.add_heading("5.1 Toast notifications (success)", level=2)
    add_table(
        doc,
        HEADERS,
        [
            R(
                "Toast",
                "Admin validate / Leave decision",
                "Conversion approved",
                "Leave conversion request approved successfully.",
                "Leave conversion approved successfully.",
            ),
            R(
                "Toast",
                "Admin validate / Leave decision",
                "Conversion rejected",
                "Leave conversion request rejected successfully.",
                "Leave conversion rejected successfully.",
            ),
            R(
                "Toast",
                "Team review / Admin validate",
                "Leave approved",
                "Leave request approved successfully.",
                "Keep (standard pattern).",
            ),
            R(
                "Toast",
                "Team review / Admin validate",
                "Leave rejected",
                "Leave request rejected successfully.",
                "Keep.",
            ),
            R(
                "Toast",
                "Admin leave requests",
                "Bulk reminder sent",
                "Reminder emails sent successfully to approvers with pending leave request.",
                '⚠ FLAG: grammar (singular). Suggested: “Reminder emails sent to approvers with pending leave requests.”',
            ),
            R(
                "Toast",
                "Admin leave request view",
                "Single reminder sent",
                "Reminder email sent successfully to the selected approver.",
                "Reminder email sent to the selected approver.",
            ),
            R(
                "Toast",
                "Admin cancel confirmed; Employee cancel pending",
                "Cancel success",
                "Leave request cancelled successfully.",
                "Keep. ⚠ Same wording for employee pending cancel and admin confirmed cancel — OK if intentional.",
            ),
            R(
                "Toast",
                "Adjust credits modal",
                "Credits adjusted",
                "Leave credits for {employeeName} updated successfully.",
                "Keep.",
            ),
            R(
                "Toast",
                "Annual distribution review",
                "Apply annual credits",
                "Annual leave credits distributed successfully to the selected employees.",
                "Annual leave credits distributed to the selected employees.",
            ),
            R(
                "Toast",
                "Contractual distribution review",
                "Apply contractual credits",
                "Leave credits distributed successfully to the selected contractual employees.",
                "Leave credits distributed to the selected contractual employees.",
            ),
            R(
                "Toast",
                "Newly regularized review",
                "Apply newly-regularized credits",
                "Leave credits distributed successfully to the selected newly regularized employees.",
                "Leave credits distributed to the selected newly regularized employees.",
            ),
            R(
                "Toast",
                "File leave (self)",
                "Submit OK",
                "Leave request submitted successfully.",
                "Keep.",
            ),
            R(
                "Toast",
                "Edit leave request",
                "Save OK",
                "Leave request updated successfully.",
                "Keep.",
            ),
            R(
                "Toast",
                "Convert to SL / Convert to EL",
                "Conversion submit OK",
                "Leave conversion submitted successfully.",
                "Keep. (Identical on both pages — good consistency.)",
            ),
            R(
                "Toast",
                "File leave for member (PIC)",
                "Filed without attachment",
                "Leave request filed for member successfully.",
                "Leave request filed for the selected member successfully.",
            ),
            R(
                "Toast",
                "File leave for member (PIC)",
                "Filed with attachment",
                "Leave request filed for member. Admin will review the attachment.",
                "Leave request filed for the selected member. An admin will review the attachment.",
            ),
        ],
    )

    # --- 5.2 Banners ---
    doc.add_heading("5.2 Banners / inline alerts", level=2)
    add_table(
        doc,
        HEADERS,
        [
            R(
                "Banner / Warning",
                "File leave (VL)",
                "Approaching quarterly VL cap",
                "Approaching quarterly Vacation Leave limit",
                "Keep.",
            ),
            R(
                "Banner / Warning",
                "File leave (VL)",
                "At quarterly VL cap",
                "Quarterly Vacation Leave limit reached",
                "Keep.",
            ),
            R(
                "Banner / Warning",
                "File leave (VL)",
                "Approaching (no projected days)",
                "You have used {used} of {limit} Vacation Leave allowed this quarter{period}. You have {remaining} remaining before the quarterly cap. This notice is informational only and does not block filing.",
                "Slightly long — consider shortening while keeping the informational disclaimer.",
            ),
            R(
                "Banner / Warning",
                "File leave (VL)",
                "At limit / with projected",
                "…Including this request, that would be {projected}. / …you would use {projected}…",
                '⚠ FLAG: inconsistent phrasing. Standardize to: “Including this request, you would use {projected}.”',
            ),
            R(
                "Banner / Alert",
                "File leave / File for member",
                "Zero credits (blocked path)",
                "Title: No remaining {type} credits · Body: You cannot file {type} without available credits.",
                'Keep body. Optional title: “No {type} credits available”.',
            ),
            R(
                "Banner / Alert",
                "File leave / File for member",
                "Duration > credits",
                "Title: Insufficient {type} credits · Body: Duration exceeds remaining credits. Shorten the leave, or file remaining days as Vacation Leave.",
                "Keep. Clear and actionable.",
            ),
            R(
                "Banner / Info",
                "File leave / File for member",
                "Full LWOP path",
                "Title: No remaining {type} credits · Body: This request will be filed as Leave Without Pay (LWOP).",
                '⚠ FLAG: same title as blocking banner. Suggested title: “Filing as Leave Without Pay (LWOP)”.',
            ),
            R(
                "Banner / Warning",
                "File leave for member only",
                "May be LWOP (PIC)",
                "Title: Insufficient leave credits · Body: Duration ({duration}) exceeds remaining {remaining}. Excess days may be treated as LWOP.",
                "Align with self-file wording where possible (“Insufficient {type} credits”).",
            ),
            R(
                "Banner / Alert",
                "File / Edit leave (VL lead time)",
                "Rule fails",
                "Filing not allowed",
                'Keep or: “Advance notice not met”.',
            ),
            R(
                "Banner",
                "File / Edit leave (VL lead time)",
                "Exact 3 workdays rule",
                "Advance filing required (3 working days)",
                "Keep.",
            ),
            R(
                "Banner",
                "File / Edit leave (VL lead time)",
                ">3 workdays rule",
                "Advance filing required (30 calendar days)",
                "Keep.",
            ),
            R(
                "Banner",
                "File / Edit leave (VL lead time)",
                "Exact-3 body",
                "Leave requests of 3 consecutive working days must be filed at least 3 working days before the leave start date.",
                "Keep (policy language).",
            ),
            R(
                "Banner",
                "File / Edit leave (VL lead time)",
                "Long leave body",
                "Leave requests of more than 3 consecutive working days must be filed at least 30 calendar days before the leave start date.",
                "Keep.",
            ),
            R(
                "Banner",
                "File / Edit leave (VL lead time)",
                "Dates OK / fail suffix",
                "Your selected dates meet this requirement. / Choose a later start date or shorten the leave duration.",
                "Keep.",
            ),
            R(
                "Banner",
                "Admin leave requests",
                "Reminder summary",
                "Pending Approver Reminders",
                "Pending approver reminders (or keep Title Case if treated as a section heading).",
            ),
            R(
                "Banner",
                "Admin leave requests",
                "Reminder loading / fail / stats",
                "Loading reminder summary… / Unable to load reminder summary. / {n} approver(s) have {m} leave request(s) awaiting their action. / Last Reminder Sent: {Never|datetime}",
                'Add “Please try again.” on load fail. Sentence-case “Last reminder sent:”.',
            ),
            R(
                "Banner / Helper",
                "Adjust credits modal",
                "Always shown (+ Adjust tooltip)",
                "For corrections and exceptions only. Use Distribute Leave Credits for routine grants.",
                "Keep (good consistency with tooltip).",
            ),
            R(
                "Banner / Warning",
                "Adjust credits modal",
                "Large ML/PL grant",
                "This is a large {Maternity|Paternity} Leave grant. Please confirm it is intentional.",
                "Keep.",
            ),
            R(
                "Banner / Info",
                "Leave conversion panels",
                "Conversion rejected",
                "This leave remains recorded as {originalLeaveType}.",
                "Keep.",
            ),
            R(
                "Banner / Alert",
                "Convert to EL",
                "EL credits < duration",
                "Remaining Emergency Leave credits ({el}) are less than this leave duration ({duration}). Conversion requires enough EL to cover the full duration, including unpaid days.",
                "OK for policy clarity; optional light trim for length.",
            ),
        ],
    )

    # --- 5.3 Confirmations ---
    doc.add_heading("5.3 Confirmation dialogs / modals", level=2)
    add_table(
        doc,
        HEADERS,
        [
            R(
                "Confirmation",
                "ConfirmReasonModal (default)",
                "Default description / empty reason / controls",
                "Please provide a reason to continue. / Reason is required. / Reason · Enter your reason… · Confirm · Go back",
                "Keep as shared defaults.",
            ),
            R(
                "Confirmation",
                "File leave",
                "Submit confirm",
                "Title: Confirm leave request · Are you sure you want to submit this leave request?",
                "Keep.",
            ),
            R(
                "Confirmation",
                "File leave for member",
                "Submit confirm",
                "Title: Confirm leave request · Submit this leave request on behalf of {name}?",
                "Keep. (Appropriately different from self-file.)",
            ),
            R(
                "Confirmation",
                "Edit leave",
                "Save confirm",
                "Title: Confirm leave request edit · Are you sure you want to save these changes to the leave request?",
                "Title: Confirm leave request changes · Are you sure you want to save these changes?",
            ),
            R(
                "Confirmation",
                "Convert SL / Convert EL",
                "Convert confirm",
                "Title: Confirm leave conversion · Are you sure you want to convert this confirmed leave request?",
                "Keep on both pages. Optional: name target type (…to Sick Leave? / …to Emergency Leave?).",
            ),
            R(
                "Confirmation",
                "Team review",
                "Approve",
                "Title: Approve Leave Request · Are you sure you want to approve this leave request?",
                '⚠ FLAG: Title Case vs Admin sentence case. Suggested title: “Approve leave request”.',
            ),
            R(
                "Confirmation",
                "Admin validate",
                "Approve leave / conversion",
                "Approve leave request / Approve leave conversion (+ matching Are you sure… bodies)",
                "Keep sentence-case titles as standard.",
            ),
            R(
                "Confirmation",
                "Team review",
                "Reject",
                "Title: Reject Leave Request · Please provide a reason… · Placeholder: Why are you rejecting this request?",
                '⚠ FLAG: Title Case + custom placeholder vs Admin. Align to “Reject leave request” and “Enter rejection reason…”.',
            ),
            R(
                "Confirmation",
                "Admin validate",
                "Reject leave / conversion",
                "Reject leave request / Reject leave conversion + Please provide a reason…",
                "Keep as standard.",
            ),
            R(
                "Confirmation",
                "Employee leave details",
                "Cancel pending",
                "Cancel leave request · This will cancel the leave request. Please provide a cancellation reason.",
                "Keep.",
            ),
            R(
                "Confirmation",
                "Admin leave view",
                "Cancel confirmed",
                "Cancel leave request · This will cancel the confirmed leave request. Please provide a cancellation reason.",
                "Keep distinction (“confirmed”). Good near-duplicate.",
            ),
            R(
                "Confirmation",
                "Cancel modals (shared)",
                "Reason UI",
                "Cancellation reason / Enter cancellation reason... / Cancel Request / Keep Request",
                "Enter cancellation reason… (consistent ellipsis). Button labels OK.",
            ),
            R(
                "Confirmation",
                "Admin leave requests / view",
                "Send reminders",
                "Send Reminder to All / Send Reminder + matching bodies",
                "Sentence-case titles: Send reminder to all / Send reminder.",
            ),
            R(
                "Confirmation",
                "Annual distribution review",
                "Apply",
                "Title: Confirm Distribution · Apply this annual leave credit distribution?",
                '⚠ FLAG: “Confirm Distribution” vs “Confirm distribution”. Standardize sentence case.',
            ),
            R(
                "Confirmation",
                "Newly regularized / Contractual review",
                "Apply",
                "Confirm distribution · Apply this newly regularized… / Apply these contractual…",
                "Keep sentence-case title; bodies are appropriately specific.",
            ),
            R(
                "Confirmation / Warning",
                "Annual & newly-regularized reviews",
                "Batch already-granted risk",
                "If anyone in this selection was already granted this year, the server will reject the whole batch. Continue only if all selected employees are still pending.",
                '⚠ FLAG: “the server will reject” is technical. Suggested: “…the entire batch will be rejected…”',
            ),
            R(
                "Modal / Helper",
                "Request conversion modal",
                "Choose SL vs EL",
                "Convert Vacation Leave · Select the leave type… · Sick/Emergency option descriptions",
                "Keep.",
            ),
            R(
                "Modal / Helper",
                "Select distribution modal",
                "Chooser",
                "Select Leave Credit Distribution · Choose how leave credits should be distributed. (+ option descriptions)",
                "Keep.",
            ),
            R(
                "Modal",
                "Adjust credits",
                "Title",
                "Adjust leave credits",
                "Keep.",
            ),
        ],
    )

    # --- 5.4 Validation ---
    doc.add_heading("5.4 Validation messages", level=2)
    add_table(
        doc,
        HEADERS,
        [
            R(
                "Validation",
                "Login",
                "Empty / invalid employee ID",
                "Please enter your employee ID. / Please enter a valid employee ID.",
                "Keep.",
            ),
            R(
                "Error",
                "Auth / Login",
                "404 not found",
                "Employee ID not found. Please check and try again.",
                "Keep.",
            ),
            R(
                "Error",
                "Auth / Login",
                "Network failure",
                "Unable to reach the server. Is the API running on port 8000?",
                '⚠ FLAG: too technical. Suggested: “Unable to reach the server. Please try again later…”',
            ),
            R(
                "Helper",
                "Login page",
                "Under field",
                "Enter your employee ID to continue.",
                "Keep.",
            ),
            R(
                "Validation",
                "Leave forms (shared)",
                "Missing fields",
                "Please select a leave type / day type / qualifying event / relationship / start date / end date.",
                "Keep (standard Please select… pattern).",
            ),
            R(
                "Validation",
                "Leave forms / ConfirmReasonModal",
                "Empty reason",
                "Reason is required.",
                "Keep.",
            ),
            R(
                "Validation",
                "Leave forms",
                "Missing attachment",
                "{Supporting document | Paternity Notification Form} is required for this leave type.",
                "Keep.",
            ),
            R(
                "Validation",
                "File leave for member",
                "Missing group/employee",
                "Please select a group. / Please select an employee.",
                "Keep.",
            ),
            R(
                "Error fallback",
                "Leave form / conversion submit",
                "API fail",
                "Unable to submit leave/conversion request. Please try again.",
                "Keep (good standard).",
            ),
            R(
                "Validation",
                "Conversion forms",
                "Reason / attachments",
                "Conversion reason is required. / At least one medical certificate… / At least one supporting document…",
                "Keep (parallel SL/EL wording).",
            ),
            R(
                "Validation",
                "Adjust credits",
                "Amount / balance / reason",
                "Credits must be greater than 0 in 0.5 steps. / Deduction cannot bring the balance below zero. / Reason must be at least 10 characters.",
                "Credits must be greater than 0, in 0.5 increments. Others: Keep.",
            ),
            R(
                "Validation",
                "File upload",
                "Duplicate / max / skip / fail",
                "That file is already attached. / You can attach up to {n} files. / Only {n} files are allowed; {skipped} skipped. / Upload failed. Please try again.",
                "Only {n} files are allowed. {skipped} file(s) were not added. Others: Keep.",
            ),
            R(
                "Helper",
                "Attachments / Dropzone",
                "Hints",
                "PDF, JPG, or PNG. Up to 5 files. / Drag & drop your file(s) here or choose file(s)",
                "Drag and drop files here, or choose files.",
            ),
            R(
                "Validation / Instruction",
                "Admin reports",
                "Year / employees not selected",
                "Select a year. / Select a conversion year. / Select a year, then click Generate Report… / Select at least one employee…",
                'Align short validations with “Please select…”. Instructional empty states: Keep.',
            ),
        ],
    )

    # --- 5.5 Errors ---
    doc.add_heading("5.5 Error / load-failure messages", level=2)
    add_table(
        doc,
        HEADERS,
        [
            R(
                "Error",
                "Leave actions",
                "Save / cancel / approve / reject / reminder / distribute",
                "Unable to [action]. Please try again. (multiple pages)",
                "Keep this pattern as the app-wide standard.",
            ),
            R(
                "Error",
                "LWOP / Monetization / Attendance",
                "Excel export fail",
                "Unable to export Excel. Please try again.",
                "Unable to export the Excel file. Please try again. (triplicate — centralize later)",
            ),
            R(
                "Error",
                "Monetization / LWOP load",
                "Load fail",
                "Unable to load monetization candidates. / Unable to load employees with LWOP.",
                'Add “Please try again.” for consistency.',
            ),
            R(
                "Error",
                "Lists / calendar / activity / PIC",
                "Load fail",
                "Unable to load leave requests / activity / upcoming leaves / members / team calendar…",
                "Keep; ensure all include Please try again.",
            ),
            R(
                "Error",
                "Admin balances / distributions / log",
                "Load fail",
                "Could not load leave balances. Try again. / Could not load eligible employees… / Could not load contractual employees… / Could not load selected employees. Go back and try again. / Could not load administrative log…",
                '⚠ FLAG: INCONSISTENT. Standardize to “Unable to load [X]. Please try again.”',
            ),
            R(
                "Error",
                "Attachment preview / Axios / attachments API",
                "Load / generic / keep-one rule",
                "Unable to load attachment… / Request failed ({status}) / Unable to update attachments. Keep at least one file.",
                'Request failed → “Something went wrong ({status}). Please try again.” Others: Keep.',
            ),
        ],
    )

    # --- 5.6 Empty ---
    doc.add_heading("5.6 Empty-state messages", level=2)
    add_table(
        doc,
        HEADERS,
        [
            R(
                "Empty-state",
                "Shared EmptyState",
                "no-data / no-results defaults",
                "No data to show / No results found",
                '⚠ “No data to show” is vague on leave lists — prefer context-specific overrides. “No results found”: Keep.',
            ),
            R(
                "Empty-state",
                "Team calendar",
                "Today / upcoming",
                "No employees are on leave today. / No upcoming approved leaves for your team.",
                "Keep.",
            ),
            R(
                "Empty-state / Helper",
                "File leave for member",
                "No members / no PIC groups",
                "No members you can file leave for in this group. / You are not assigned to any PIC groups.",
                "Keep.",
            ),
            R(
                "Empty-state",
                "Employee multi-select / FileField / Preview",
                "Various",
                "No employees in group / No employees match your search. / No file chosen / No attachment to preview.",
                "No employees in this group. Others: Keep.",
            ),
        ],
    )

    # --- 5.7 Tooltips ---
    doc.add_heading("5.7 Tooltips & disabled-state titles", level=2)
    add_table(
        doc,
        HEADERS,
        [
            R(
                "Tooltip",
                "Admin balances / Paternity / Forms",
                "Adjust / download / uploading / no changes",
                "For corrections and exceptions only… / Download the Paternity Notification Form… / Wait for the attachment upload to finish. / No changes to save",
                "Keep.",
            ),
            R(
                "Tooltip",
                "Convert pages",
                "Disabled confirm / EL credits",
                "Enter a conversion reason… / Upload at least one medical certificate… / supporting document… / Checking Emergency Leave credits. / Remaining Emergency Leave credits are less than…",
                "Keep (parallel SL/EL wording).",
            ),
            R(
                "Tooltip",
                "Details / Team / Admin",
                "Edit/cancel disabled",
                "Only pending leave requests can be cancelled / Editing is not available during leave conversion / Editing is disabled after an approver has taken action / Leave with a pending conversion cannot be cancelled…",
                "Add trailing periods for consistency. Conversion cancel guidance: Keep.",
            ),
            R(
                "Tooltip",
                "Lists / Monetization / Preview / Header",
                "Misc",
                "Already granted this year / At least one file must remain / Remove file / Not applicable / Download unavailable — no file URL / Loading file… / Expand/Collapse sidebar",
                "Keep; light punctuation cleanup optional.",
            ),
        ],
    )

    # --- 5.8 Status ---
    doc.add_heading("5.8 Status-related messages & labels", level=2)
    add_table(
        doc,
        HEADERS,
        [
            R(
                "Status badge",
                "Leave status system",
                "List / detail badges",
                "PENDING APPROVAL / PENDING CONVERSION / APPROVED / CONVERTED TO SL / CONVERTED TO EL / REJECTED / CANCELLED",
                "Keep ALL CAPS if intentional visual system.",
            ),
            R(
                "Status filter",
                "Dashboard / lists",
                "Filter dropdown",
                "All statuses / Pending Approval / Pending conversion / Approved / Converted to SL / Converted to EL / Rejected / Cancelled",
                '⚠ FLAG: “Pending conversion” casing. Suggested: “Pending Conversion”.',
            ),
            R(
                "Status",
                "Admin list / Approver badge",
                "Overrides / approval table",
                "Pending Conversion / PENDING / PENDING REVIEW / Approved / Rejected / No Action Yet / No action required",
                '⚠ FLAG: “No Action Yet” vs “No action required”. Align casing.',
            ),
        ],
    )

    # --- 5.9 Error pages ---
    doc.add_heading("5.9 System / error pages", level=2)
    add_table(
        doc,
        HEADERS,
        [
            R(
                "Error page",
                "404 / Leave not found / 403 / 401 / 409 / 500 / 503 / Offline",
                "Full-page errors",
                "Page not found / Leave request not found / You do not have access… / Your session has expired / This request has already been updated / Something went wrong / The service is temporarily unavailable / You are offline (+ bodies)",
                "Keep — clear and user-friendly overall.",
            ),
            R(
                "Inline error",
                "Details / review / validate / convert",
                "Not found / conversion unavailable",
                "No leave request was found for ID {id}… / Conversion not available — Only your own confirmed Vacation Leave…",
                "Prefer less emphasis on raw ID for end users. SL/EL parallel wording: Keep.",
            ),
            R(
                "Helper",
                "Edit leave",
                "Type locked",
                "Leave type cannot be changed after filing.",
                "Keep.",
            ),
        ],
    )

    # --- 5.10 Placeholders ---
    doc.add_heading("5.10 Instructional placeholders (selected)", level=2)
    add_table(
        doc,
        HEADERS,
        [
            R(
                "Placeholder",
                "Leave forms",
                "Selects / reason",
                "Select Leave Type / Day Type / Qualifying Event / Relationship / Group / Employee / Enter reason for leave",
                "Optional: sentence-case selects to match validation tone.",
            ),
            R(
                "Placeholder",
                "Conversion / Cancel / Reject / Adjust",
                "Reason fields",
                "Enter reason for converting to Sick/Emergency Leave / Enter cancellation reason... / Why are you rejecting this request? / Enter reason for this adjustment (min. 10 characters)…",
                'Align reject placeholder with Admin: “Enter rejection reason…”.',
            ),
            R(
                "Placeholder",
                "Search fields",
                "Lists",
                "Search employee name or number / Search by name or number",
                '⚠ Near-duplicate. Standardize: “Search by name or employee number”.',
            ),
        ],
    )

    # --- 5.11 Attachment ---
    doc.add_heading("5.11 Attachment preview messages", level=2)
    add_table(
        doc,
        HEADERS,
        [
            R(
                "Empty / Status / Error",
                "Attachment preview",
                "No file / loading / no URL",
                "No attachment to preview. / Loading attachment… / No file URL is available for this attachment.",
                "Keep.",
            ),
        ],
    )

    doc.add_heading("6. Duplicate & inconsistency checklist", level=1)
    for d in [
        "Reminder bulk toast: singular “leave request” vs plural meaning.",
        "Approve/Reject modal titles: Team uses Title Case; Admin uses sentence case.",
        "Reject placeholder: Team custom vs Admin default.",
        "Confirm Distribution vs Confirm distribution.",
        "Pending conversion vs Pending Conversion vs PENDING CONVERSION.",
        "Unable to… Please try again. vs Could not… Try again.",
        "No remaining credits title used for both block and LWOP info.",
        "Quarterly VL projected phrasing: “that would be” vs “you would use”.",
        "Search placeholders differ slightly across pages.",
        "Approver badges: “No Action Yet” vs “No action required”.",
        "Excel export error identical on 3 pages (centralize when implementing).",
        "Conversion submit toast / confirm identical on SL & EL (good consistency).",
    ]:
        bullet(doc, d)

    doc.add_heading("7. Next steps", level=1)
    for i, s in enumerate(
        [
            "Review the Suggested Revision column; mark Accept / Edit / Keep as-is.",
            "Agree on standards in Section 4 (especially error fallbacks and title casing).",
            "Decide whether badges stay ALL CAPS while filters use Title Case.",
            "After sign-off, implement via shared constants (extend toastMessages.ts pattern) — no code changed in this review.",
        ],
        1,
    ):
        p = doc.add_paragraph()
        r = p.add_run(f"{i}. {s}")
        set_run_font(r, size=10)

    p = doc.add_paragraph()
    r = p.add_run(
        "Source: leaveforms-frontend/src (pages, components, constants, utils, auth)."
    )
    set_run_font(r, size=9, color=RGBColor(89, 89, 89))

    doc.save(OUT_PATH)
    print(f"Wrote {OUT_PATH} ({os.path.getsize(OUT_PATH)} bytes)")


if __name__ == "__main__":
    main()
